"""
LocalMindDesk — PPT 制作工具
根据 LLM 生成的大纲自动创建精美 PPT
"""
import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from app import llm_provider

# PPT 主题色
THEME = {
    "bg": RGBColor(0x0F, 0x17, 0x2A),       # 深蓝背景
    "title_bg": RGBColor(0x1A, 0x25, 0x40),  # 标题页背景
    "accent": RGBColor(0x3B, 0x82, 0xF6),    # 蓝色强调
    "text": RGBColor(0xFA, 0xFA, 0xFA),      # 白色文字
    "sub": RGBColor(0xA1, 0xA1, 0xAA),       # 灰色副文字
    "card": RGBColor(0x1E, 0x29, 0x3B),      # 卡片背景
}

def _set_slide_bg(slide, color):
    """设置幻灯片背景色"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def _add_text_box(slide, left, top, width, height, text, font_size=18,
                  color=None, bold=False, alignment=PP_ALIGN.LEFT):
    """添加文本框"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color or THEME["text"]
    p.font.bold = bold
    p.alignment = alignment
    return txBox

def create_title_slide(prs, title, subtitle=""):
    """创建标题页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    _set_slide_bg(slide, THEME["title_bg"])

    # 标题
    _add_text_box(slide, Inches(1), Inches(2.5), Inches(8), Inches(1.2),
                  title, font_size=36, bold=True, alignment=PP_ALIGN.CENTER)

    # 副标题
    if subtitle:
        _add_text_box(slide, Inches(1), Inches(3.8), Inches(8), Inches(0.8),
                      subtitle, font_size=16, color=THEME["sub"], alignment=PP_ALIGN.CENTER)

    # 底部装饰线
    from pptx.util import Emu
    shape = slide.shapes.add_shape(
        1, Inches(3), Inches(4.8), Inches(4), Pt(3)  # 矩形
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = THEME["accent"]
    shape.line.fill.background()

def create_content_slide(prs, title, bullets):
    """创建内容页：标题 + 要点列表"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, THEME["bg"])

    # 页面标题
    _add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8),
                  title, font_size=28, bold=True, color=THEME["accent"])

    # 分割线
    shape = slide.shapes.add_shape(1, Inches(0.8), Inches(1.2), Inches(8.4), Pt(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = THEME["accent"]
    shape.line.fill.background()

    # 内容要点
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = f"  •  {bullet}"
        p.font.size = Pt(16)
        p.font.color.rgb = THEME["text"]
        p.space_after = Pt(12)
        p.line_spacing = Pt(24)

def create_end_slide(prs, text="谢谢！"):
    """创建结尾页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, THEME["title_bg"])
    _add_text_box(slide, Inches(1), Inches(3), Inches(8), Inches(1.5),
                  text, font_size=40, bold=True, alignment=PP_ALIGN.CENTER,
                  color=THEME["accent"])

def generate_ppt_outline(topic: str) -> dict:
    """用 LLM 生成 PPT 大纲（JSON 格式）"""
    prompt = f"""请为以下主题生成 PPT 大纲，直接输出 JSON，不要输出其他内容。

主题：{topic}

要求的 JSON 格式：
{{
  "title": "PPT 总标题",
  "subtitle": "副标题",
  "slides": [
    {{
      "title": "第一页标题",
      "bullets": ["要点1", "要点2", "要点3"]
    }},
    {{
      "title": "第二页标题",
      "bullets": ["要点1", "要点2", "要点3"]
    }}
  ]
}}

要求：
- 5-8 页内容（不含封面和结尾页）
- 每页 3-5 个要点
- 内容专业、有条理
- 直接输出 JSON，不要 markdown 代码块"""

    reply = llm_provider.chat(
        messages=[{"role": "user", "content": prompt}],
        system_prompt="你是一个专业的 PPT 制作助手。严格按要求输出 JSON。",
        temperature=0.3,
    )

    # 清理可能的 markdown 代码块包裹
    reply = reply.strip()
    if reply.startswith("```"):
        lines = reply.split("\n")
        reply = "\n".join(lines[1:-1] if lines[-1].strip() == "```" else lines[1:])
    reply = reply.strip()

    try:
        return json.loads(reply)
    except json.JSONDecodeError:
        # 尝试提取 JSON 片段
        start = reply.find("{")
        end = reply.rfind("}") + 1
        if start >= 0 and end > start:
            return json.loads(reply[start:end])
        raise ValueError(f"LLM 返回的不是有效 JSON：{reply[:200]}")

def make_ppt(topic: str, outline: dict = None) -> str:
    """
    生成 PPT 文件，返回文件路径
    如果 outline 为空，自动用 LLM 生成大纲
    """
    if outline is None:
        outline = generate_ppt_outline(topic)

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 封面
    create_title_slide(prs, outline.get("title", topic), outline.get("subtitle", ""))

    # 内容页
    for slide_data in outline.get("slides", []):
        title = slide_data.get("title", "")
        bullets = slide_data.get("bullets", [])
        if title:
            create_content_slide(prs, title, bullets)

    # 结尾
    create_end_slide(prs)

    # 保存
    os.makedirs("data/files", exist_ok=True)
    safe_name = "".join(c for c in topic[:30] if c.isalnum() or c in " _-").strip() or "presentation"
    filepath = f"data/files/{safe_name}.pptx"
    prs.save(filepath)

    print(f"[PPT] 已生成: {filepath} ({len(outline.get('slides', []))} 页)")
    return filepath
