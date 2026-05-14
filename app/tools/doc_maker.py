"""
LocalMindDesk — Word 文档生成工具
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def make_docx(title: str, content: str) -> str:
    """
    从 Markdown 文本生成 Word 文档
    content: Markdown 格式文本
    返回文件路径
    """
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)

    # 标题
    doc.add_heading(title, level=0)

    # 解析 Markdown 内容
    lines = content.split('\n')
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif stripped[0:3] in ('1. ', '2. ', '3. ', '4. ', '5. ', '6. ', '7. ', '8. ', '9. '):
            doc.add_paragraph(stripped[3:], style='List Number')
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            run = p.add_run(stripped[2:])
            run.italic = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif stripped.startswith('```'):
            continue  # 跳过代码块标记
        else:
            # 处理加粗
            if '**' in stripped:
                p = doc.add_paragraph()
                parts = stripped.split('**')
                for i, part in enumerate(parts):
                    if part:
                        run = p.add_run(part)
                        if i % 2 == 1:
                            run.bold = True
            else:
                doc.add_paragraph(stripped)

    # 保存
    os.makedirs("data/files", exist_ok=True)
    safe_name = "".join(c for c in title[:30] if c.isalnum() or c in " _-").strip() or "document"
    filepath = f"data/files/{safe_name}.docx"
    doc.save(filepath)

    print(f"[DOC] 已生成: {filepath}")
    return filepath
